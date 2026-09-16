"""
bravo_doc_helpers.py
====================
Módulo Python com todos os helpers para construir documentos longos Bravo
no padrão consulting-grade.

Uso típico:

    from bravo_doc_helpers import (
        setup_document, add_heading1, add_heading2, add_heading3,
        add_lead, add_body, add_bullet, add_callout, add_quote,
        add_bravo_table, add_code_block, add_kpi_strip,
        attach_header, attach_footer, insert_cover_image,
        BRAVO,
    )

    doc, content_width_dxa, content_width_cm = setup_document()
    insert_cover_image(doc, '/path/to/cover_full.png')
    new_section = open_content_section(doc)
    attach_header(new_section, '/path/to/bravo-logo-color-alpha.png', 'Título do doc · Área')
    attach_footer(new_section, 'Confidencial — Uso interno Bravo')

    add_heading1(doc, '1.  Capítulo um', new_page=False)
    add_lead(doc, 'Parágrafo de abertura do capítulo.')
    ...
    doc.save('saida.docx')

Todas as funções aqui implementam EXATAMENTE a lógica validada na versão final
do documento 'Fundamentos de IA Aplicada' (maio de 2026). Não improvise — se
precisar de algo novo, adicione um helper aqui e documente.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ════════════════════════════════════════════════════════════════════════
#                            PALETA E TOKENS
# ════════════════════════════════════════════════════════════════════════
class BRAVO:
    """Paleta Bravo + tokens tipográficos consolidados para documentos."""

    # Cores como RGBColor (para python-docx)
    PURPLE       = RGBColor(0x30, 0x2D, 0x71)
    PURPLE_DARK  = RGBColor(0x20, 0x1D, 0x49)
    PURPLE_DEEP  = RGBColor(0x3D, 0x35, 0x9E)
    PURPLE_MID   = RGBColor(0x45, 0x3D, 0x84)
    JACARTA      = RGBColor(0x80, 0x62, 0xDE)
    AMETHYST     = RGBColor(0xAA, 0x6C, 0xC9)
    VIKING       = RGBColor(0x5E, 0xCA, 0xDA)
    LIGHT_PUR    = RGBColor(0xAA, 0x97, 0xEF)
    LIGHTER      = RGBColor(0xC8, 0xBF, 0xF3)
    BG_LIGHT     = RGBColor(0xEC, 0xEA, 0xFA)
    WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
    GREY_MID     = RGBColor(0x66, 0x66, 0x66)
    GREY_DARK    = RGBColor(0x1A, 0x1A, 0x2E)
    GREEN_POS    = RGBColor(0xB6, 0xD7, 0xA8)
    GREEN_TXT    = RGBColor(0x2D, 0x6E, 0x1E)
    YELLOW_NEU   = RGBColor(0xFF, 0xF2, 0xCC)
    YELLOW_TXT   = RGBColor(0x7A, 0x5F, 0x00)
    RED_NEG      = RGBColor(0xF4, 0xCC, 0xCC)
    RED_TXT      = RGBColor(0x9B, 0x1C, 0x1C)

    # Mesmas cores como hex string (para shading/borders XML)
    PURPLE_H       = "302D71"
    PURPLE_DARK_H  = "201D49"
    PURPLE_DEEP_H  = "3D359E"
    PURPLE_MID_H   = "453D84"
    JACARTA_H      = "8062DE"
    AMETHYST_H     = "AA6CC9"
    VIKING_H       = "5ECADA"
    LIGHT_PUR_H    = "AA97EF"
    LIGHTER_H      = "C8BFF3"
    BG_LIGHT_H     = "ECEAFA"
    BG_LIGHTER_H   = "F5F3FB"
    WHITE_H        = "FFFFFF"
    GREEN_POS_H    = "B6D7A8"
    GREEN_TXT_H    = "2D6E1E"
    YELLOW_NEU_H   = "FFF2CC"
    YELLOW_TXT_H   = "7A5F00"
    RED_NEG_H      = "F4CCCC"
    RED_TXT_H      = "9B1C1C"

    # Tipografia
    FONT       = "Montserrat"
    FONT_MONO  = "Consolas"


# ════════════════════════════════════════════════════════════════════════
#                       PRIMITIVAS XML DE BAIXO NÍVEL
# ════════════════════════════════════════════════════════════════════════

def set_cell_shading(cell, hex_color):
    """Define cor de fundo de célula via w:shd CLEAR (nunca SOLID)."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    for existing in tcPr.findall(qn('w:shd')):
        tcPr.remove(existing)
    tcPr.append(shd)


def set_cell_borders(cell, color="FFFFFF", size="6",
                     sides=("top", "left", "bottom", "right")):
    """size em oitavos de pt (6 = 0.75pt; 36 = 4.5pt para acentos)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for s in sides:
        b = tcBorders.find(qn(f'w:{s}'))
        if b is None:
            b = OxmlElement(f'w:{s}')
            tcBorders.append(b)
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), size)
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), color)


def set_cell_margins(cell, top=80, left=140, bottom=80, right=140):
    """Padding interno em twips (1pt = 20 twips)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn('w:tcMar'))
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for side, val in (('top', top), ('left', left),
                      ('bottom', bottom), ('right', right)):
        node = tcMar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tcMar.append(node)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')


def set_table_width(table, dxa_total):
    """Fixa largura total da tabela + layout fixed (não autofit)."""
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(dxa_total))
    tblW.set(qn('w:type'), 'dxa')

    tblLayout = tblPr.find(qn('w:tblLayout'))
    if tblLayout is None:
        tblLayout = OxmlElement('w:tblLayout')
        tblPr.append(tblLayout)
    tblLayout.set(qn('w:type'), 'fixed')


def style_run(run, *, font=BRAVO.FONT, size=11, bold=False, italic=False, color=None):
    """Aplica fonte + tamanho + cor garantindo que rFonts está completo
    (ascii, hAnsi, cs, eastAsia)."""
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(k), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


# ════════════════════════════════════════════════════════════════════════
#                       SETUP DE DOCUMENTO E SEÇÕES
# ════════════════════════════════════════════════════════════════════════

def setup_document():
    """Cria Document base configurado para capa full-bleed.

    A PRIMEIRA section é configurada com margens zero (para capa em imagem A4).
    Depois você chama open_content_section(doc) para abrir a section de conteúdo
    com margens normais.

    Retorna: (doc, content_width_dxa, content_width_cm)
    """
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)

    # Margens zero na capa
    section.top_margin = Cm(0)
    section.bottom_margin = Cm(0)
    section.left_margin = Cm(0)
    section.right_margin = Cm(0)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)

    # titlePg: esconde header/footer na primeira página
    sectPr = section._sectPr
    titlePg = sectPr.find(qn('w:titlePg'))
    if titlePg is None:
        titlePg = OxmlElement('w:titlePg')
        sectPr.append(titlePg)

    # Estilo Normal global
    normal = doc.styles['Normal']
    normal.font.name = BRAVO.FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = BRAVO.GREY_DARK
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(k), BRAVO.FONT)

    # Content width: 210mm - 2*22mm = 166mm = 16.6cm
    content_width_cm = 16.6
    content_width_dxa = int(content_width_cm * 567)  # 1cm ≈ 567 dxa
    return doc, content_width_dxa, content_width_cm


def insert_cover_image(doc, cover_png_path):
    """Insere imagem de capa A4 inteira (210x297mm) na primeira section."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(cover_png_path, width=Mm(210), height=Mm(297))


def open_content_section(doc, *, top=2.4, bottom=2.0, left=2.2, right=2.2):
    """Abre nova section (NEW_PAGE) com margens normais e header/footer
    desvinculados da seção anterior (capa).

    Retorna a section criada — passe para attach_header() e attach_footer().
    """
    new_section = doc.add_section(WD_SECTION.NEW_PAGE)
    new_section.page_width = Mm(210)
    new_section.page_height = Mm(297)
    new_section.top_margin = Cm(top)
    new_section.bottom_margin = Cm(bottom)
    new_section.left_margin = Cm(left)
    new_section.right_margin = Cm(right)
    new_section.header_distance = Cm(1.0)
    new_section.footer_distance = Cm(0.9)
    new_section.header.is_linked_to_previous = False
    new_section.footer.is_linked_to_previous = False
    return new_section


# ════════════════════════════════════════════════════════════════════════
#                          HEADER E FOOTER
# ════════════════════════════════════════════════════════════════════════

def attach_header(section, logo_alpha_png_path, doc_title_right,
                  content_width_dxa=None, content_width_cm=16.6):
    """Header da página 2+: tabela invisível 2 colunas (logo | título à direita).

    logo_alpha_png_path: PATH para o logo colorido com transparência alfa real
                         (use o arquivo bravo-logo-color-alpha.png desta skill).
    doc_title_right:     string ex: "Título do Doc · Bravo AI Operations"
    """
    if content_width_dxa is None:
        content_width_dxa = int(content_width_cm * 567)

    header = section.header
    header.is_linked_to_previous = False
    for p in list(header.paragraphs):
        p._element.getparent().remove(p._element)

    tbl = header.add_table(rows=1, cols=2, width=Cm(content_width_cm))
    tbl.autofit = False
    set_table_width(tbl, content_width_dxa)

    tblGrid = tbl._tbl.find(qn('w:tblGrid'))
    if tblGrid is None:
        tblGrid = OxmlElement('w:tblGrid')
        tbl._tbl.insert(0, tblGrid)
    for gc in list(tblGrid):
        tblGrid.remove(gc)
    for w in [int(5.0 * 567), int((content_width_cm - 5.0) * 567)]:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(w))
        tblGrid.append(gc)

    c1 = tbl.rows[0].cells[0]; c1.width = Cm(5.0)
    c2 = tbl.rows[0].cells[1]; c2.width = Cm(content_width_cm - 5.0)
    for c in (c1, c2):
        set_cell_margins(c, top=0, left=0, bottom=0, right=0)
        set_cell_borders(c, color="FFFFFF", size="4")
        set_cell_shading(c, "FFFFFF")

    # Logo à esquerda
    c1.paragraphs[0].text = ''
    pl = c1.paragraphs[0]
    pl.paragraph_format.space_after = Pt(0)
    pl.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rl = pl.add_run()
    rl.add_picture(logo_alpha_png_path, height=Cm(0.85))

    # Título à direita
    c2.paragraphs[0].text = ''
    pt_ = c2.paragraphs[0]
    pt_.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pt_.paragraph_format.space_after = Pt(0)
    rt = pt_.add_run(doc_title_right)
    style_run(rt, size=9, color=BRAVO.PURPLE_MID)

    # Linha separadora fina
    div = header.add_paragraph()
    div.paragraph_format.space_before = Pt(2)
    div.paragraph_format.space_after = Pt(0)
    pPr = div._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single')
    bot.set(qn('w:sz'), '4')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), BRAVO.LIGHTER_H)
    pBdr.append(bot); pPr.append(pBdr)


def attach_footer(section, classification_left,
                  content_width_dxa=None, content_width_cm=16.6):
    """Footer: classificação à esquerda + 'Pág. X' dinâmico à direita."""
    if content_width_dxa is None:
        content_width_dxa = int(content_width_cm * 567)

    footer = section.footer
    footer.is_linked_to_previous = False
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)

    tbl = footer.add_table(rows=1, cols=2, width=Cm(content_width_cm))
    tbl.autofit = False
    set_table_width(tbl, content_width_dxa)

    tblGrid = tbl._tbl.find(qn('w:tblGrid'))
    if tblGrid is None:
        tblGrid = OxmlElement('w:tblGrid')
        tbl._tbl.insert(0, tblGrid)
    for gc in list(tblGrid):
        tblGrid.remove(gc)
    for w in [int(11.0 * 567), int((content_width_cm - 11.0) * 567)]:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(w))
        tblGrid.append(gc)

    c1 = tbl.rows[0].cells[0]; c1.width = Cm(11.0)
    c2 = tbl.rows[0].cells[1]; c2.width = Cm(content_width_cm - 11.0)
    for c in (c1, c2):
        set_cell_margins(c, top=80, left=0, bottom=0, right=0)
        set_cell_borders(c, color="FFFFFF", size="4")
        set_cell_shading(c, "FFFFFF")

    # Classificação à esquerda
    c1.paragraphs[0].text = ''
    pf1 = c1.paragraphs[0]
    pf1.paragraph_format.space_after = Pt(0)
    rf1 = pf1.add_run(classification_left)
    style_run(rf1, size=8.5, color=BRAVO.GREY_MID)

    # Paginação à direita (field PAGE)
    c2.paragraphs[0].text = ''
    pf2 = c2.paragraphs[0]
    pf2.paragraph_format.space_after = Pt(0)
    pf2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rpre = pf2.add_run("Pág. ")
    style_run(rpre, size=8.5, color=BRAVO.GREY_MID)

    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    fld_inst = OxmlElement('w:instrText')
    fld_inst.set(qn('xml:space'), 'preserve')
    fld_inst.text = ' PAGE '
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    rfield = pf2.add_run()
    style_run(rfield, size=8.5, color=BRAVO.GREY_MID)
    rfield._element.append(fld_begin)
    rfield._element.append(fld_inst)
    rfield._element.append(fld_end)

    # Linha separadora em cima do footer
    div = footer.add_paragraph()
    div.paragraph_format.space_after = Pt(0)
    pPr = div._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single')
    top.set(qn('w:sz'), '4')
    top.set(qn('w:space'), '1')
    top.set(qn('w:color'), BRAVO.LIGHTER_H)
    pBdr.append(top); pPr.append(pBdr)


# ════════════════════════════════════════════════════════════════════════
#                       HEADINGS, LEADS E PARÁGRAFOS
# ════════════════════════════════════════════════════════════════════════

def add_heading1(doc, txt, *, new_page=True):
    """H1: 24pt Bold Purple + borda inferior Purple grossa. Nova página por padrão."""
    if new_page:
        p_break = doc.add_paragraph()
        r = p_break.add_run()
        r.add_break(WD_BREAK.PAGE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(txt)
    style_run(r, size=24, bold=True, color=BRAVO.PURPLE)

    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '18')
    bottom.set(qn('w:space'), '6')
    bottom.set(qn('w:color'), BRAVO.PURPLE_H)
    pBdr.append(bottom)
    for existing in pPr.findall(qn('w:pBdr')):
        pPr.remove(existing)
    pPr.append(pBdr)

    spacing = doc.add_paragraph()
    spacing.paragraph_format.space_after = Pt(6)


def add_heading2(doc, txt):
    """H2: 16pt Bold Purple, 20pt antes, 8pt depois, keep_with_next."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(txt)
    style_run(r, size=16, bold=True, color=BRAVO.PURPLE)


def add_heading3(doc, txt):
    """H3: 13pt Bold Purple Mid, 14pt antes, 4pt depois."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(txt)
    style_run(r, size=13, bold=True, color=BRAVO.PURPLE_MID)


def add_lead(doc, text):
    """Lead: 12pt Purple Mid, line-spacing 1.5 — primeiro parágrafo de capítulo."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    style_run(r, size=12, color=BRAVO.PURPLE_MID)


def add_body(doc, runs):
    """Corpo de parágrafo com inline styling.
    runs: lista de (texto, {bold:, italic:, color:, size:}).
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.5
    for txt, attrs in runs:
        r = p.add_run(txt)
        style_run(r,
                  size=attrs.get('size', 11),
                  bold=attrs.get('bold', False),
                  italic=attrs.get('italic', False),
                  color=attrs.get('color', BRAVO.GREY_DARK))
    return p


def add_bullet(doc, runs, level=0):
    """Bullet visual com marcador '•' Jacarta bold. Inline styling como add_body."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.left_indent = Cm(0.6 + level * 0.8)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    r = p.add_run("•  ")
    style_run(r, size=11, bold=True, color=BRAVO.JACARTA)
    for txt, attrs in runs:
        r = p.add_run(txt)
        style_run(r,
                  size=attrs.get('size', 11),
                  bold=attrs.get('bold', False),
                  italic=attrs.get('italic', False),
                  color=attrs.get('color', BRAVO.GREY_DARK))


def add_numbered(doc, idx, text):
    """Item numerado simples: '1.  Texto.' Purple bold + texto Grey Dark."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    r = p.add_run(f"{idx}.  ")
    style_run(r, size=11, bold=True, color=BRAVO.PURPLE)
    r2 = p.add_run(text)
    style_run(r2, size=11, color=BRAVO.GREY_DARK)


# ════════════════════════════════════════════════════════════════════════
#                              CALLOUTS
# ════════════════════════════════════════════════════════════════════════

_CALLOUT_PALETTE = {
    # kind: (bg_hex, accent_hex, body_text_hex)
    'info':      (BRAVO.BG_LIGHT_H,  BRAVO.JACARTA_H,    BRAVO.PURPLE_H),
    'rule':      (BRAVO.BG_LIGHT_H,  BRAVO.JACARTA_H,    BRAVO.PURPLE_H),
    'objective': (BRAVO.BG_LIGHT_H,  BRAVO.JACARTA_H,    BRAVO.PURPLE_H),
    'warn':      ("FFF8DD",          BRAVO.YELLOW_TXT_H, BRAVO.YELLOW_TXT_H),
    'danger':    ("F9E4E4",          BRAVO.RED_TXT_H,    BRAVO.RED_TXT_H),
    'success':   ("E8F4E3",          BRAVO.GREEN_TXT_H,  BRAVO.GREEN_TXT_H),
}


def add_callout(doc, label, text, *, kind='info',
                content_width_dxa=None, content_width_cm=16.6):
    """Callout 1×1 com border-left grossa. label em maiúsculas 9pt bold."""
    if content_width_dxa is None:
        content_width_dxa = int(content_width_cm * 567)
    bg, accent, txt_col = _CALLOUT_PALETTE.get(kind, _CALLOUT_PALETTE['info'])

    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table, content_width_dxa)

    cell = table.rows[0].cells[0]
    cell.width = Cm(content_width_cm)
    set_cell_shading(cell, bg)
    set_cell_margins(cell, top=180, left=240, bottom=180, right=240)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, sz, col in (('top', '4', bg), ('right', '4', bg), ('bottom', '4', bg)):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), sz)
        b.set(qn('w:space'), '0');    b.set(qn('w:color'), col)
        tcBorders.append(b)
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), '36')
    left.set(qn('w:space'), '0');    left.set(qn('w:color'), accent)
    tcBorders.append(left)
    for existing in tcPr.findall(qn('w:tcBorders')):
        tcPr.remove(existing)
    tcPr.append(tcBorders)

    # Label
    cell.paragraphs[0].text = ''
    p_lbl = cell.paragraphs[0]
    p_lbl.paragraph_format.space_after = Pt(4)
    rl = p_lbl.add_run(label.upper())
    style_run(rl, size=9, bold=True, color=RGBColor.from_string(accent))

    # Corpo
    p_body = cell.add_paragraph()
    p_body.paragraph_format.space_after = Pt(0)
    p_body.paragraph_format.line_spacing = 1.4
    rb = p_body.add_run(text)
    style_run(rb, size=11, color=RGBColor.from_string(txt_col))

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(8)


def add_quote(doc, text, source=None,
              content_width_dxa=None, content_width_cm=16.6):
    """Quote em destaque: border-left Jacarta + texto 13pt italic Purple."""
    if content_width_dxa is None:
        content_width_dxa = int(content_width_cm * 567)

    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, content_width_dxa)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "FFFFFF")
    set_cell_margins(cell, top=120, left=300, bottom=120, right=240)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side, sz, col in (('top', '4', 'FFFFFF'),
                          ('right', '4', 'FFFFFF'),
                          ('bottom', '4', 'FFFFFF')):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), sz)
        b.set(qn('w:space'), '0');    b.set(qn('w:color'), col)
        tcBorders.append(b)
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), '36')
    left.set(qn('w:space'), '0');    left.set(qn('w:color'), BRAVO.JACARTA_H)
    tcBorders.append(left)
    for existing in tcPr.findall(qn('w:tcBorders')):
        tcPr.remove(existing)
    tcPr.append(tcBorders)

    cell.paragraphs[0].text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.45
    r = p.add_run(f"“{text}”")
    style_run(r, size=13, italic=True, color=BRAVO.PURPLE)
    if source:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(4)
        r2 = p2.add_run(f"— {source}")
        style_run(r2, size=10, color=BRAVO.GREY_MID)

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(8)


# ════════════════════════════════════════════════════════════════════════
#                              TABELAS
# ════════════════════════════════════════════════════════════════════════

def add_bravo_table(doc, headers, rows, *, col_widths_cm=None,
                    header_color=BRAVO.PURPLE_H, total_row=None,
                    footer_caption=None, unit=None,
                    content_width_cm=16.6):
    """Tabela on-brand Bravo com largura fixa, zebra ECEAFA, cabeçalho Purple.

    rows: lista de listas. Cada célula pode ser:
        - string simples, OU
        - dict {"text": ..., "bold": bool, "italic": bool, "color": RGBColor, "size": int}
    total_row: opcional, linha extra com fundo Jacarta + texto branco bold.
    unit: opcional, string que aparece à direita acima da tabela (ex: "(R$ MM)").
    footer_caption: opcional, "Tabela X. Descrição."
    """
    n = len(headers)
    if col_widths_cm is None:
        each = content_width_cm / n
        col_widths_cm = [each] * n
    else:
        # ajusta se diferenças de arredondamento
        diff = content_width_cm - sum(col_widths_cm)
        if abs(diff) < 0.05:
            col_widths_cm[-1] += diff

    # Unit annotation
    if unit:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cap.paragraph_format.space_after = Pt(2)
        cr = cap.add_run(unit)
        style_run(cr, size=9, italic=True, color=BRAVO.GREY_MID)

    n_rows_total = 1 + len(rows) + (1 if total_row else 0)
    table = doc.add_table(rows=n_rows_total, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    dxa_widths = [int(w * 567) for w in col_widths_cm]
    set_table_width(table, sum(dxa_widths))

    # tblGrid explícito
    tblGrid = table._tbl.find(qn('w:tblGrid'))
    if tblGrid is None:
        tblGrid = OxmlElement('w:tblGrid')
        table._tbl.insert(0, tblGrid)
    for gc in list(tblGrid):
        tblGrid.remove(gc)
    for w in dxa_widths:
        gc = OxmlElement('w:gridCol')
        gc.set(qn('w:w'), str(w))
        tblGrid.append(gc)

    # Header
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        c = hdr[i]
        c.width = Cm(col_widths_cm[i])
        set_cell_shading(c, header_color)
        set_cell_borders(c, color="FFFFFF", size="6")
        set_cell_margins(c, top=160, left=160, bottom=160, right=160)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        c.paragraphs[0].text = ''
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(h)
        style_run(r, size=10.5, bold=True, color=BRAVO.WHITE)

    # Linhas de dado
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        zebra = (r_idx % 2 == 1)
        for i, val in enumerate(row):
            c = cells[i]
            c.width = Cm(col_widths_cm[i])
            set_cell_shading(c, BRAVO.BG_LIGHT_H if zebra else "FFFFFF")
            set_cell_borders(c, color="FFFFFF", size="6")
            set_cell_margins(c, top=120, left=160, bottom=120, right=160)
            c.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            c.paragraphs[0].text = ''
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.35
            if isinstance(val, dict):
                r = p.add_run(val.get('text', ''))
                style_run(r,
                          size=val.get('size', 10),
                          bold=val.get('bold', False),
                          italic=val.get('italic', False),
                          color=val.get('color', BRAVO.PURPLE_MID))
            else:
                r = p.add_run(str(val))
                style_run(r, size=10, color=BRAVO.PURPLE_MID)

    # Linha total
    if total_row:
        cells = table.rows[-1].cells
        for i, val in enumerate(total_row):
            c = cells[i]
            c.width = Cm(col_widths_cm[i])
            set_cell_shading(c, BRAVO.JACARTA_H)
            set_cell_borders(c, color="FFFFFF", size="6")
            set_cell_margins(c, top=140, left=160, bottom=140, right=160)
            c.paragraphs[0].text = ''
            p = c.paragraphs[0]
            r = p.add_run(str(val))
            style_run(r, size=11, bold=True, color=BRAVO.WHITE)

    # Caption
    if footer_caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(4)
        cap.paragraph_format.space_after = Pt(10)
        cr = cap.add_run(footer_caption)
        style_run(cr, size=9, italic=True, color=BRAVO.GREY_MID)
    else:
        after = doc.add_paragraph()
        after.paragraph_format.space_after = Pt(6)


# ════════════════════════════════════════════════════════════════════════
#                       CODE BLOCKS E KPI STRIPS
# ════════════════════════════════════════════════════════════════════════

def add_code_block(doc, text, *, label=None,
                   content_width_dxa=None, content_width_cm=16.6):
    """Code/prompt block: fundo F5F3FB, fonte Consolas 9.5pt, label opcional."""
    if content_width_dxa is None:
        content_width_dxa = int(content_width_cm * 567)

    if label:
        p_l = doc.add_paragraph()
        p_l.paragraph_format.space_after = Pt(2)
        rl = p_l.add_run(label)
        style_run(rl, size=9, bold=True, color=BRAVO.JACARTA)

    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, content_width_dxa)
    c = table.rows[0].cells[0]
    set_cell_shading(c, BRAVO.BG_LIGHTER_H)
    set_cell_borders(c, color=BRAVO.BG_LIGHT_H, size="6")
    set_cell_margins(c, top=160, left=200, bottom=160, right=200)

    c.paragraphs[0].text = ''
    for i, ln in enumerate(text.split('\n')):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.35
        r = p.add_run(ln if ln else ' ')
        style_run(r, font=BRAVO.FONT_MONO, size=9.5, color=BRAVO.PURPLE)

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(6)


def add_kpi_strip(doc, kpis,
                  content_width_dxa=None, content_width_cm=16.6):
    """KPI strip horizontal. kpis: lista de (numero_str, label_str). 3-4 itens."""
    if content_width_dxa is None:
        content_width_dxa = int(content_width_cm * 567)
    n = len(kpis)
    table = doc.add_table(rows=1, cols=n)
    set_table_width(table, content_width_dxa)
    each = content_width_cm / n
    for i, (num, lbl) in enumerate(kpis):
        c = table.rows[0].cells[i]
        c.width = Cm(each)
        set_cell_shading(c, BRAVO.BG_LIGHT_H)
        set_cell_margins(c, top=240, left=80, bottom=240, right=80)
        tcPr = c._tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        for side, sz, col in (('top', '12', BRAVO.JACARTA_H),
                              ('right', '4', BRAVO.WHITE_H),
                              ('bottom', '4', BRAVO.BG_LIGHT_H),
                              ('left', '4', BRAVO.WHITE_H)):
            b = OxmlElement(f'w:{side}')
            b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), sz)
            b.set(qn('w:space'), '0');    b.set(qn('w:color'), col)
            tcBorders.append(b)
        for existing in tcPr.findall(qn('w:tcBorders')):
            tcPr.remove(existing)
        tcPr.append(tcBorders)

        c.paragraphs[0].text = ''
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        rn = p.add_run(num)
        style_run(rn, size=28, bold=True, color=BRAVO.PURPLE)

        p2 = c.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        rl = p2.add_run(lbl)
        style_run(rl, size=9, color=BRAVO.PURPLE_MID)

    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(10)


# ════════════════════════════════════════════════════════════════════════
#                       SUMÁRIO/ÍNDICE MANUAL
# ════════════════════════════════════════════════════════════════════════

def add_toc_entry(doc, num, title, subtitles=()):
    """Adiciona uma entrada principal de índice + subtítulos.
    subtitles: tuple/list de strings (cada uma vira linha indentada)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    rn = p.add_run(f"{num}   ")
    style_run(rn, size=12, bold=True, color=BRAVO.JACARTA)
    rt = p.add_run(title)
    style_run(rt, size=12, bold=True, color=BRAVO.PURPLE)
    for s in subtitles:
        ps = doc.add_paragraph()
        ps.paragraph_format.space_after = Pt(2)
        ps.paragraph_format.line_spacing = 1.3
        ps.paragraph_format.left_indent = Cm(1.0)
        rs = ps.add_run(s)
        style_run(rs, size=10.5, color=BRAVO.PURPLE_MID)


# ════════════════════════════════════════════════════════════════════════
#                       FECHAMENTO / COLOPHON
# ════════════════════════════════════════════════════════════════════════

def add_closing_colophon(doc, text="— FIM DO DOCUMENTO —"):
    """Linha centralizada discreta no fim do documento."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    style_run(r, size=9, bold=True, color=BRAVO.LIGHT_PUR)
